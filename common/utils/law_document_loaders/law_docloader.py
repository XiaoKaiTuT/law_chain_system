# 该脚本用于加载Word文档

from typing import Iterator
from .law_ocr import get_ocr
from tqdm import tqdm                   # 进度条库，处理大量文件时显示进度，例如：[████▒▒▒▒] 45%
from docx.table import _Cell, Table                 # Word里的表格和表格里的单元格
from docx.oxml.table import CT_Tbl                  # Word底层的XML表格节点，用于精确遍历文档结构
from docx.oxml.text.paragraph import CT_P           # Word底层的XML段落节点，用于精确遍历文档结构
from docx.text.paragraph import Paragraph           # Word里的一段文字/段落
from docx import Document as Docu1                  # 打开一个Word文件(.docx)
from docx.document import Document as Docu2         # Word文档的完整对象类型，用来标注变量类型
from docx import ImagePart                          # 从Word文件里提取嵌入的图片
from PIL import Image                               # 打开、处理图片(裁剪、缩放等)，PIL是Python最经典的图片处理库
from io import BytesIO                              # 把内存中的二进制数据当文件用(比如从Word里取出的图片数据，不用先存到硬盘就能直接处理)
import numpy as np
from langchain_core.documents import Document
from langchain_core.document_loaders import BaseLoader      # LangChain的加载器基类

# OCR文档加载器，用于加载Word文档
class OCRDOCLoader(BaseLoader):
    def __init__(self, filepath: str) -> None:
        """
        函数功能：初始化OCR文档加载器
        :param filepath: 文件路径
        :return: None
        """
        self.filepath = filepath  # 文件路径


    def lazy_load(self) -> Iterator[Document]:
        """
        函数功能：惰性加载文档，返回一个生成器，每次返回一个文档
        :return: 文档生成器
        """
        line = self.doc2text(self.filepath)
        yield Document(page_content=line, metadata={"source": self.filepath})

    def doc2text(self, filepath):
        """
        函数功能：将Word文档转换为文本
        :param filepath: 文件路径
        :return: 文本内容
        """
        # 创建OCR识别对象
        ocr = get_ocr()
        # 读取word文档
        doc = Docu1(filepath)
        # 定义一个空字符串，用于存储最终文本内容
        resp = ""

        def iter_block_items(parent):
            """
            函数作用：定义一个迭代器，用于遍历文档中的块(段落、表格等)
            :param parent: 文档块Document对象
            :return: 文档块对象
            """
            # 判断parent对象类型，如果是Document类型，则获取其内容
            if isinstance(parent, Docu2):
                parent_elm = parent.element.body
            # 如果是表格单元格类型，获取单元格的xml元素
            elif isinstance(parent, _Cell):
                parent_elm = parent._tc
            else:
                raise ValueError("文档解析失败")
            # 遍历parent_elm中的所有子元素
            for child in parent_elm.iterchildren():
                # 如果是段落类型，yield段落对象
                if isinstance(child, CT_P):
                    yield Paragraph(child, parent)
                # 如果是表格类型，yield表格对象
                elif isinstance(child, CT_Tbl):
                    yield Table(child, parent)

        # 创建进度条，表示文档处理的进度
        # 参1：进度条总长度(段落数+表格数)         参2：进度条描述
        b_unit = tqdm(total=len(doc.paragraphs) + len(doc.tables), desc="OCRDOCLoader block index: 0")
        # 遍历文档中的所有块(段落和表格)
        for i, block in enumerate(iter_block_items(doc)):
            b_unit.set_description("OCRDOCLoader block index: {}".format(i))
            b_unit.refresh()        # 刷新进度条

            # 如果块是段落类型
            if isinstance(block, Paragraph):
                resp += block.text.strip() + "\n"       # 将段落文本加入到返回字符串中
                # 获取段落中的所有图片
                images = block._element.xpath('.//pic:pic')
                for image in images:
                    # 遍历图片，获取图片ID
                    for img_id in image.xpath('.//a:blip/@r.embed'):
                        part = doc.part.related_parts[img_id]       # 根据图片ID获取图片对象
                        # 如果该部分是图片
                        if isinstance(part, ImagePart):
                            image = Image.open(BytesIO(part._blob))     # 打开图片
                            result, _ = ocr(np.array(image))        # 使用OCR识别图片中的文字
                            # 如果识别结果不为空
                            if result:
                                ocr_result = [line[1] for line in result]   # 提取识别出的文字
                                resp += "\n".join(ocr_result)       # 将识别结果加入返回文本中
            # 如果块是表格类型
            elif isinstance(block, Table):
                # 遍历表格中的所有行和单元格
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            resp += paragraph.text.strip() + "\n"       # 将单元格内的段落文本加入返回文本中

            # 更新进度条
            b_unit.update(1)
        # 返回提取的文本内容
        return resp

# 测试函数
if __name__ == '__main__':
    docx_loader = OCRDOCLoader(
        r"../test_data/中华人民共和国民事诉讼法（《关于修改〈中华人民共和国民事诉讼法〉的决定》第五次修正））.docx")
    doc = docx_loader.load()
    # [Document(metadata={source:"..."}, page_content="..."), ...]
    print(doc)